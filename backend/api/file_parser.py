"""Parse uploaded files to extract text content for AI analysis."""

import json
import csv
import io
from pathlib import Path


def parse_file(filepath: str, filename: str) -> dict:
    """
    Parse a file and return its text content.
    Returns {"text": "...", "parse_status": "ok"|"empty"|"unsupported"}
    """
    ext = Path(filename).suffix.lower()
    parser = PARSERS.get(ext)
    if not parser:
        return {"text": "", "parse_status": "unsupported"}
    try:
        text = parser(filepath)
        if text and text.strip():
            return {"text": text.strip()[:50000], "parse_status": "ok"}
        return {"text": "", "parse_status": "empty"}
    except Exception:
        return {"text": "", "parse_status": "empty"}


# ── Plain text formats ──

def _read_text(filepath: str) -> str:
    for enc in ("utf-8", "gbk", "latin-1"):
        try:
            with open(filepath, "r", encoding=enc) as f:
                return f.read()
        except (UnicodeDecodeError, LookupError):
            continue
    return ""


def _parse_json(filepath: str) -> str:
    for enc in ("utf-8", "gbk", "latin-1"):
        try:
            with open(filepath, "r", encoding=enc) as f:
                data = json.load(f)
            return json.dumps(data, ensure_ascii=False, indent=2)
        except (json.JSONDecodeError, UnicodeDecodeError, LookupError):
            continue
    return ""


def _parse_csv(filepath: str) -> str:
    for enc in ("utf-8", "gbk", "latin-1"):
        try:
            with open(filepath, "r", encoding=enc, newline="") as f:
                reader = csv.reader(f)
                rows = list(reader)
            if not rows:
                return ""
            return "\n".join(" | ".join(row) for row in rows)
        except (UnicodeDecodeError, LookupError):
            continue
    return ""


# ── PDF ──

def _parse_pdf(filepath: str) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        return ""
    reader = PdfReader(filepath)
    pages = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            pages.append(text)
    return "\n\n".join(pages)


# ── DOCX ──

def _parse_docx(filepath: str) -> str:
    try:
        from docx import Document
    except ImportError:
        return ""
    doc = Document(filepath)
    parts = []
    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text)
    for table in doc.tables:
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            rows.append(" | ".join(cells))
        if rows:
            parts.append("\n".join(rows))
    return "\n\n".join(parts)


# ── XLSX ──

def _parse_xlsx(filepath: str) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return ""
    wb = load_workbook(filepath, data_only=True)
    all_sheets = []
    for name in wb.sheetnames:
        ws = wb[name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            if any(c for c in cells):
                rows.append(" | ".join(cells))
        if rows:
            all_sheets.append(f"## Sheet: {name}\n" + "\n".join(rows))
    return "\n\n".join(all_sheets)


PARSERS = {
    ".txt": _read_text,
    ".md": _read_text,
    ".log": _read_text,
    ".xml": _read_text,
    ".json": _parse_json,
    ".csv": _parse_csv,
    ".pdf": _parse_pdf,
    ".docx": _parse_docx,
    ".xlsx": _parse_xlsx,
}
