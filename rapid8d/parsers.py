"""Local text extraction with explicit coverage, never OCR by implication."""
import io,zipfile
from pathlib import Path
from .domain import UserError
MAX_FILE=20*1024*1024
MAX_TEXT=120000
def extract(name,raw):
    if len(raw)>MAX_FILE: raise UserError("单个文件请控制在20 MB以内。")
    suffix=Path(name).suffix.lower(); warnings=[]; chunks=[]; used=0
    def add(location,text):
        nonlocal used
        text=str(text).strip()
        if not text: return
        room=MAX_TEXT-used
        if room<=0: return
        fragment=text[:room]; chunks.append(dict(locator=location,text=fragment)); used+=len(fragment)
        if len(fragment)<len(text) and "文字提取达到120000字符上限，存在未提取内容。" not in warnings:
            warnings.append("文字提取达到120000字符上限，存在未提取内容。")
    if suffix in (".docx",".xlsx"):
        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if sum(x.file_size for x in archive.infolist())>160*1024*1024:
                    raise UserError("文件展开后过大，请拆分后导入。")
        except zipfile.BadZipFile: raise UserError("文件格式损坏，无法读取。")
    if suffix in (".txt",".md",".csv",".tsv"):
        for encoding in ("utf-8-sig","gb18030","utf-16"):
            try: text=raw.decode(encoding); break
            except UnicodeError: continue
        else: raise UserError("无法识别文本编码，请另存为UTF-8。")
        for i,line in enumerate(text.splitlines(),1): add(f"行 {i}",line)
        coverage="文本行；CSV/TSV按原行保留"
    elif suffix==".docx":
        from docx import Document
        from docx.text.paragraph import Paragraph
        from docx.table import Table
        doc=Document(io.BytesIO(raw)); pi=ti=0
        for element in doc.element.body:
            if element.tag.endswith("}p"):
                pi+=1; add(f"正文段落 {pi}",Paragraph(element,doc).text)
            elif element.tag.endswith("}tbl"):
                ti+=1
                for ri,row in enumerate(Table(element,doc).rows,1):
                    add(f"表 {ti} 行 {ri}"," | ".join(c.text for c in row.cells))
        warnings.append("已读取正文和表格；未识别图片、批注、文本框和页眉页脚。")
        coverage=f"正文 {pi} 段、{ti} 张表"
    elif suffix==".pdf":
        from pypdf import PdfReader
        reader=PdfReader(io.BytesIO(raw))
        if reader.is_encrypted and not reader.decrypt(""): raise UserError("PDF已加密，请先提供可读取的副本。")
        if len(reader.pages)>200: raise UserError("PDF超过200页，请拆分相关页后导入。")
        empty=[]
        for i,page in enumerate(reader.pages,1):
            try: text=page.extract_text() or ""
            except Exception: text=""
            if not text.strip(): empty.append(i)
            add(f"第 {i} 页",text)
        if empty: warnings.append("以下页未提取到文字（可能为扫描件）："+",".join(map(str,empty)))
        warnings.append("PDF按页面提取文字；图片和复杂表格的布局未作视觉识别。")
        coverage=f"检查了 {len(reader.pages)} 页；{len(empty)} 页无可读文字"
    elif suffix==".xlsx":
        from openpyxl import load_workbook
        wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
        formulas=load_workbook(io.BytesIO(raw),read_only=True,data_only=False)
        try:
            for ws in wb:
                if ws.max_row>5000 or ws.max_column>80: warnings.append(f"{ws.title}只提取前5000行/80列，其余未读取。")
                for row in ws.iter_rows(min_row=1,max_row=min(ws.max_row,5000),max_col=min(ws.max_column,80)):
                    values=[f"{c.coordinate}={c.value}" for c in row if c.value is not None]
                    if values: add(ws.title+"!"+str(row[0].row)," | ".join(values))
                fw=formulas[ws.title]
                count=sum(1 for row in fw.iter_rows(max_row=min(fw.max_row,5000),max_col=min(fw.max_column,80)) for c in row if c.data_type=="f")
                if count: warnings.append(f"{ws.title}含{count}个公式；使用文件缓存值，不重新计算，请在Excel重算保存后核对。")
            coverage=f"{len(wb.sheetnames)} 个工作表（含隐藏表）；单元格带地址"
            warnings.append("未解析图表、图片、批注或宏。")
        finally: wb.close(); formulas.close()
    elif suffix==".xls":
        import xlrd
        wb=xlrd.open_workbook(file_contents=raw)
        for ws in wb.sheets():
            if ws.nrows>5000 or ws.ncols>80: warnings.append(f"{ws.name}只提取前5000行/80列。")
            for r in range(min(ws.nrows,5000)):
                add(f"{ws.name} 行 {r+1}"," | ".join(str(x) for x in ws.row_values(r,0,min(ws.ncols,80))))
        coverage=f"{wb.nsheets} 个工作表；单元格缓存值"
        warnings.append("旧版Excel按行提取值，未重算公式或解析图片。")
    else: raise UserError("支持 DOCX、文字PDF、XLSX、XLS、CSV、TXT、MD。旧版DOC请另存为DOCX；照片与扫描件识别暂未启用。")
    if used>=MAX_TEXT and not any("120000" in w for w in warnings): warnings.append("已达到120000字符上限，可能有后续内容未提取。")
    if not chunks: raise UserError("没有提取到可用文字。扫描件/图片暂不支持自动识别，请粘贴文字或上传文字版文件。")
    return dict(chunks=chunks,coverage=coverage,warnings=warnings,char_count=used)
